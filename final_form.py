"""
Batch VGGT + L2CS + DeepGaze IIE pipeline — FINAL FORM

Changes from my_ending_data__with_GradCAM_subfolder.py
- Grad-CAM fully reformed:
    * GradCAMSingleBin removed; replaced by GradCAM context manager.
    * compute_weighted_gradcam_for_layer_topk removed; replaced by compute_gradcam_for_layer.
    * One forward pass + at most two backward passes (one per head) instead of
      iterating N_bins backward passes per layer.
    * Hook registration/cleanup is guaranteed via __enter__/__exit__.
    * _norm01 extracted as a module-level helper.
- Everything else is identical to the original.
"""

from __future__ import annotations

import os
import json
import csv
import warnings
from glob import glob
from typing import Dict, List, Optional, Tuple, Any

import gc
import re
import openpyxl
from openpyxl.styles import Font, PatternFill
import cv2
import numpy as np
import torch
import torch.nn as nn
from PIL import Image
from scipy.special import logsumexp

from l2cs import Pipeline  # type: ignore
from vggt.models.vggt import VGGT  # type: ignore
from vggt.utils.load_fn import load_and_preprocess_images  # type: ignore
from vggt.utils.pose_enc import pose_encoding_to_extri_intri  # type: ignore
from blob_projection_utils import project_blob_to_images  # type: ignore
import deepgaze_pytorch  # type: ignore

warnings.filterwarnings("ignore", message=".*flash attention.*", category=UserWarning)

# =============================================================================
# CONFIG
# =============================================================================
# Point to the top-level experiment folder — everything else is auto-discovered.
BASE_ROOT = r"/home/keisokulab/Downloads/Jikken_2 (3rd copy)"

# Optional filters — set None to process everything found under BASE_ROOT.
# ONLY_SCENARIO_NAME   : name of the scenario root folder  (e.g. "scenario_1_4m")
# ONLY_SUBSCENARIO_NAME: name of the sub-scenario folder   (e.g. "testing_1scenario_right_4m")
ONLY_SCENARIO_NAME: Optional[str] = None
ONLY_SUBSCENARIO_NAME: Optional[str] = None

RESUME_AUTO = True
RESUME_SKIP_SKIPPED = False
DONE_MARKER = "_DONE.json"
SKIP_MARKER = "_SKIPPED.json"

REF_IMAGE_NAMES = ["view1.jpg", "view2.jpg"]

DEEPGAZE_VIEW_IDXS = [2]
GAZE_VIEW_IDX = 0

L2CS_WEIGHTS = r"/home/keisokulab/learning_workshop/L2CSNet_gaze360.pkl"
CENTERBIAS_NPY = r"centerbias_mit1003.npy"

DEEPGAZE_MAX_DIM = 1024

DYNAMIC_CONE = True
GAZE_CONE_ANGLE = 50.0

NUM_BINS = 90
BIN_WIDTH_DEG = 4.0
BIN_START_DEG = -180.0

CONE_QUANTILE = 0.95
CONE_SAMPLES = 2000
CONE_MIN_DEG = 8.0
CONE_MAX_DEG = 60.0

MASK_OUTSIDE_DARK = 0.0

BETA = 1.0
ALPHA = 1.0
EPS = 1e-12

SINGLEOBJ_ENABLE = True
SINGLEOBJ_REL_THRESH = 0.50
SINGLEOBJ_SMOOTH_SIGMA = 0.0
SINGLEOBJ_MIN_AREA = 50
SINGLEOBJ_MORPH_CLOSE = True

ANNOTATE_IF_MISSING = True
GT_FILENAME = "ground_truth_polygons.json"
DO_EVAL = True

USE_SOFT_GT_FOR_METRICS = True
GT_SOFT_SIGMA_MODE = "min_dim_frac"
GT_SOFT_SIGMA_VALUE = 0.03
GT_SOFT_SIGMA_MIN_PX = 1.0
GT_SOFT_SIGMA_MAX_PX = 40.0

USE_GRADCAM_L2CS = False
GRADCAM_ALPHA = 0.60
SAVE_GRADCAM_NPY = False

GRADCAM_FORCE_CPU = True
GRADCAM_LAYER_MODE = "all"
GRADCAM_SINGLE_LAYER_NAME = "layer4"
GRADCAM_TOPK_BINS = 90
GRADCAM_COMPUTE = "both"

# Cone half-angle ablation configurations.
# Each entry is run independently for every primary image; outputs go to a
# named subdirectory so results are easy to compare side-by-side.
CONE_ABLATION = [
    # Dynamic: cone size derived from the L2CS probability distribution at a given quantile.
    {"name": "dynamic_q99", "dynamic": True,  "quantile": 0.99, "fixed_deg": None},
    {"name": "dynamic_q95", "dynamic": True,  "quantile": 0.95, "fixed_deg": None},
    {"name": "dynamic_q87", "dynamic": True,  "quantile": 0.87, "fixed_deg": None},
    # Fixed: hard-coded cone half-angle.
    {"name": "fixed_60",    "dynamic": False, "quantile": None,  "fixed_deg": 60.0},
    {"name": "fixed_30",    "dynamic": False, "quantile": None,  "fixed_deg": 30.0},
    {"name": "fixed_15",    "dynamic": False, "quantile": None,  "fixed_deg": 15.0},
]

# Output folder for per-person per-distance Excel summaries (created automatically).
SUMMARY_OUT_DIR = r"/home/keisokulab/Desktop/final_form"

# Maps the top-level person folder name under BASE_ROOT to a readable label.
PERSON_LABEL_MAP: Dict[str, str] = {
    "1Senpai":  "senpai",
    "Moriyama": "moriyama",
    "Yoneyama": "yoneyama",
}

# Numeric columns included in the per-sub-scenario average row written to Excel.
SUMMARY_AVG_COLS = [
    "cone_half_angle_deg", "sigma_theta_deg",
    "auc", "ap",
    "dist_px_to_gt_centroid", "dist_px_to_gt_centroid_l2cs",
    "angular_error_l2cs_deg", "angular_error_peak_deg",
]

# =============================================================================
# DEVICE
# =============================================================================
device_str = "cuda" if torch.cuda.is_available() else "cpu"
torch_dev = torch.device(device_str)
dtype = torch.float16 if device_str == "cuda" else torch.float32

if torch.cuda.is_available():
    torch.backends.cuda.matmul.allow_tf32 = True
    torch.backends.cudnn.allow_tf32 = True


# =============================================================================
# HELPERS (geometry, fusion, utilities)
# =============================================================================
def gaze_vector_from_angles(yaw_rad: float, pitch_rad: float) -> np.ndarray:
    x = np.sin(yaw_rad) * np.cos(pitch_rad)
    y = np.sin(pitch_rad)
    z = np.cos(pitch_rad) * np.cos(yaw_rad)
    g = np.array([x, y, z], dtype=np.float32)
    return -g / (np.linalg.norm(g) + 1e-9)


def _clip_bbox_xyxy(bbox_xyxy: np.ndarray, h: int, w: int) -> Tuple[int, int, int, int]:
    x1, y1, x2, y2 = bbox_xyxy.astype(int).tolist()
    x1 = max(0, min(w - 1, x1))
    x2 = max(0, min(w, x2))
    y1 = max(0, min(h - 1, y1))
    y2 = max(0, min(h, y2))

    if x2 <= x1 + 2 or y2 <= y1 + 2:
        cx, cy = w // 2, h // 2
        r = min(w, h) // 4
        x1, x2 = max(0, cx - r), min(w, cx + r)
        y1, y2 = max(0, cy - r), min(h, cy + r)

    return x1, y1, x2, y2


def preprocess_face_for_l2cs_with_vis(
    img_bgr: np.ndarray,
    bbox_xyxy: np.ndarray,
    dev: torch.device,
    out_size: int = 224,
) -> Tuple[torch.Tensor, np.ndarray]:
    h, w = img_bgr.shape[:2]
    x1, y1, x2, y2 = _clip_bbox_xyxy(bbox_xyxy, h, w)

    face_bgr = img_bgr[y1:y2, x1:x2].copy()
    face_vis_bgr = cv2.resize(face_bgr, (out_size, out_size), interpolation=cv2.INTER_LINEAR)

    face_rgb = cv2.cvtColor(face_vis_bgr, cv2.COLOR_BGR2RGB)
    face_t = torch.from_numpy(face_rgb).permute(2, 0, 1).float() / 255.0
    mean = torch.tensor([0.485, 0.456, 0.406], dtype=torch.float32).view(3, 1, 1)
    std = torch.tensor([0.229, 0.224, 0.225], dtype=torch.float32).view(3, 1, 1)
    face_t = (face_t - mean) / std
    return face_t.unsqueeze(0).to(dev), face_vis_bgr


def expected_angle_deg(probs: np.ndarray, bin_angles_deg: np.ndarray) -> float:
    probs = probs.astype(np.float64)
    probs = probs / (probs.sum() + 1e-12)
    return float((probs * bin_angles_deg.astype(np.float64)).sum())


def cone_half_angle_from_probs(
    yaw_probs: np.ndarray,
    pitch_probs: np.ndarray,
    bin_angles_deg: np.ndarray,
    num_samples: int = 2000,
    quantile: float = 0.90,
    min_deg: float = 8.0,
    max_deg: float = 80.0,
) -> float:
    yaw_probs = yaw_probs.astype(np.float64)
    pitch_probs = pitch_probs.astype(np.float64)
    yaw_probs /= (yaw_probs.sum() + 1e-12)
    pitch_probs /= (pitch_probs.sum() + 1e-12)

    yaw_s_deg = np.random.choice(bin_angles_deg, size=num_samples, p=yaw_probs)
    pitch_s_deg = np.random.choice(bin_angles_deg, size=num_samples, p=pitch_probs)
    yaw_s = np.deg2rad(yaw_s_deg)
    pitch_s = np.deg2rad(pitch_s_deg)

    x = np.sin(yaw_s) * np.cos(pitch_s)
    y = np.sin(pitch_s)
    z = np.cos(pitch_s) * np.cos(yaw_s)
    vecs = np.stack([x, y, z], axis=1).astype(np.float32)
    vecs = -vecs / (np.linalg.norm(vecs, axis=1, keepdims=True) + 1e-9)

    mean = vecs.mean(axis=0)
    mean = mean / (np.linalg.norm(mean) + 1e-9)

    cosang = np.clip(vecs @ mean, -1.0, 1.0)
    ang_deg = np.rad2deg(np.arccos(cosang))
    half_angle = float(np.quantile(ang_deg, quantile))
    return float(np.clip(half_angle, min_deg, max_deg))


def estimate_eye_origin_world(
    bbox,
    world_points_view,
    orig_size,
    H_net,
    W_net,
    intrinsic,
    extrinsic,
):
    x1, y1, x2, y2 = map(float, bbox)
    H_i, W_i = orig_size

    cx_orig = 0.5 * (x1 + x2)
    cy_orig = 0.5 * (y1 + y2)

    u_net = cx_orig * (W_net / W_i)
    v_net = cy_orig * (H_net / H_i)

    u_i = int(np.clip(round(u_net), 0, W_net - 1))
    v_i = int(np.clip(round(v_net), 0, H_net - 1))

    point_world = world_points_view[v_i, u_i]

    if not np.isfinite(point_world).all():
        r = 2
        v0, v1 = max(0, v_i - r), min(H_net, v_i + r + 1)
        u0, u1 = max(0, u_i - r), min(W_net, u_i + r + 1)
        patch = world_points_view[v0:v1, u0:u1].reshape(-1, 3)
        finite_mask = np.isfinite(patch).all(axis=1)
        if finite_mask.any():
            point_world = np.median(patch[finite_mask], axis=0)
        else:
            return point_world.astype(np.float32), (int(cx_orig), int(cy_orig))

    R = extrinsic[:3, :3]
    t = extrinsic[:3, 3]
    point_cam = R @ point_world + t
    depth = float(point_cam[2])

    fx = float(intrinsic[0, 0])
    fy = float(intrinsic[1, 1])
    cx = float(intrinsic[0, 2])
    cy = float(intrinsic[1, 2])

    fx_orig = fx * (W_i / W_net)
    fy_orig = fy * (H_i / H_net)
    cx_orig_K = cx * (W_i / W_net)
    cy_orig_K = cy * (H_i / H_net)

    x_cam = (cx_orig - cx_orig_K) / (fx_orig + 1e-12) * depth
    y_cam = (cy_orig - cy_orig_K) / (fy_orig + 1e-12) * depth
    z_cam = depth

    point_cam_new = np.array([x_cam, y_cam, z_cam], dtype=np.float32)
    point_world_new = R.T @ (point_cam_new - t)
    return point_world_new.astype(np.float32), (int(cx_orig), int(cy_orig))


def compute_cone_mask_for_points(
    pts: np.ndarray,
    origin3d: np.ndarray,
    dir3d_world: np.ndarray,
    cone_angle_deg: float,
    min_face_dist: float = 0.20,
):
    vecs = pts - origin3d[None, :]
    dist = np.linalg.norm(vecs, axis=1)
    far_enough = dist > min_face_dist
    vecs_n = vecs / (dist[:, None] + 1e-9)
    cosang = np.clip(vecs_n @ dir3d_world, -1.0, 1.0)
    ang_deg = np.rad2deg(np.arccos(cosang))
    in_cone = ang_deg <= cone_angle_deg
    return in_cone & far_enough, dist


def _mask_to_u8(mask: np.ndarray) -> np.ndarray:
    if mask.ndim == 3:
        mask = cv2.cvtColor(mask, cv2.COLOR_BGR2GRAY)
    if mask.dtype == np.bool_:
        return (mask.astype(np.uint8) * 255)
    if mask.dtype == np.uint8:
        return mask
    m = mask.astype(np.float32)
    mx = float(np.nanmax(m)) if m.size else 0.0
    if mx <= 1.0 + 1e-6:
        m = m * 255.0
    m = np.clip(m, 0, 255)
    return m.astype(np.uint8)


def blend_image_with_mask(img_bgr: np.ndarray, mask_u8: np.ndarray, outside_dark: float = 0.20) -> np.ndarray:
    alpha = (mask_u8.astype(np.float32) / 255.0)[..., None]
    base = img_bgr.astype(np.float32)
    dark = base * float(outside_dark)
    out = base * alpha + dark * (1.0 - alpha)
    return np.clip(out, 0, 255).astype(np.uint8)


def gaze_dir_cam_from_dir_world(dir_world: np.ndarray, extrinsic_w2c: np.ndarray) -> np.ndarray:
    R_wc = extrinsic_w2c[:3, :3].astype(np.float32)
    d = (R_wc @ dir_world.astype(np.float32)).astype(np.float32)
    d /= (np.linalg.norm(d) + 1e-9)
    return d


def theta_map_from_eye_and_worldpoints(
    world_points_view: np.ndarray,
    extrinsic_w2c: np.ndarray,
    eye_world: np.ndarray,
    gaze_dir_cam: np.ndarray,
) -> np.ndarray:
    R = extrinsic_w2c[:3, :3].astype(np.float32)
    t = extrinsic_w2c[:3, 3].astype(np.float32)

    eye_cam = (R @ eye_world.astype(np.float32) + t).astype(np.float32)

    Pw = world_points_view.reshape(-1, 3).astype(np.float32)
    Pc = (R @ Pw.T).T + t[None, :]

    V = Pc - eye_cam[None, :]
    Vn = V / (np.linalg.norm(V, axis=1, keepdims=True) + 1e-9)

    g = gaze_dir_cam.astype(np.float32)
    g /= (np.linalg.norm(g) + 1e-9)

    cosang = np.clip(Vn @ g, -1.0, 1.0)
    theta = np.arccos(cosang).astype(np.float32)
    theta = theta.reshape(world_points_view.shape[0], world_points_view.shape[1])

    bad = ~np.isfinite(theta)
    if bad.any():
        theta[bad] = np.pi
    return theta


def sigma_from_cone_quantile(cone_half_angle_rad: float, quantile: float) -> float:
    q = float(np.clip(quantile, 1e-4, 1.0 - 1e-4))
    val = torch.tensor(2.0 * q - 1.0, dtype=torch.float32)
    z = float(np.sqrt(2.0) * torch.special.erfinv(val).item())
    z = max(z, 1e-6)
    return float(cone_half_angle_rad / z)


def geometry_likelihood(theta_rad: np.ndarray, sigma_theta: float) -> np.ndarray:
    return np.exp(-0.5 * (theta_rad / (sigma_theta + 1e-12)) ** 2).astype(np.float32)


def normalize_prob_map(x: np.ndarray) -> np.ndarray:
    x = x.astype(np.float32)
    s = float(np.sum(x))
    if not np.isfinite(s) or s <= 0:
        return np.full_like(x, 1.0 / x.size, dtype=np.float32)
    return x / s


def fuse_poe(G: np.ndarray, S: np.ndarray, alpha: float, beta: float, eps: float = 1e-12) -> np.ndarray:
    logH = alpha * np.log(G + eps) + beta * np.log(S + eps)
    logH -= float(logsumexp(logH))
    H = np.exp(logH).astype(np.float32)
    return normalize_prob_map(H)


def overlay_heatmap_bgr(image_bgr: np.ndarray, heat: np.ndarray, out_path: str, alpha: float = 0.6) -> None:
    heat = heat.astype(np.float32)
    heat_u8 = (255.0 * heat / (float(heat.max()) + 1e-12)).clip(0, 255).astype(np.uint8)
    heat_color = cv2.applyColorMap(heat_u8, cv2.COLORMAP_MAGMA)
    out = (image_bgr.astype(np.float32) * (1 - alpha) + heat_color.astype(np.float32) * alpha).clip(0, 255).astype(np.uint8)
    cv2.imwrite(out_path, out)


# =============================================================================
# Grad-CAM helpers (L2CS) — reformed
# =============================================================================

def find_last_conv2d(model: nn.Module) -> Tuple[str, nn.Conv2d]:
    last_name: Optional[str] = None
    last_mod: Optional[nn.Conv2d] = None
    for name, m in model.named_modules():
        if isinstance(m, nn.Conv2d):
            last_name = name
            last_mod = m
    if last_name is None or last_mod is None:
        raise RuntimeError("Could not find any nn.Conv2d layer for Grad-CAM.")
    return last_name, last_mod


def find_resnet_layers(model: nn.Module) -> Dict[str, nn.Module]:
    layers: Dict[str, nn.Module] = {}
    for name, m in model.named_modules():
        if name in ("layer1", "layer2", "layer3", "layer4"):
            layers[name] = m
        elif name.endswith(".layer1"):
            layers["layer1"] = m
        elif name.endswith(".layer2"):
            layers["layer2"] = m
        elif name.endswith(".layer3"):
            layers["layer3"] = m
        elif name.endswith(".layer4"):
            layers["layer4"] = m

    _last_name, last_conv_mod = find_last_conv2d(model)
    layers["lastconv"] = last_conv_mod
    return layers


def _norm01(cam: np.ndarray) -> np.ndarray:
    cam = cam.astype(np.float32)
    cam -= cam.min()
    mx = float(cam.max())
    return cam / mx if mx > 0 else cam


class GradCAM:
    """
    Context-manager Grad-CAM for a single target layer.

    Registers forward/backward hooks on enter and removes them on exit,
    guaranteeing no hook leaks regardless of exceptions.

    Usage:
        with GradCAM(model, layer) as gcam:
            yaw_logits, pitch_logits = model(x)
            cam = gcam.cam_from_score(weighted_score, retain_graph=True)
    """

    def __init__(self, model: nn.Module, target_layer: nn.Module) -> None:
        self.model = model
        self.target_layer = target_layer
        self._acts: Optional[torch.Tensor] = None
        self._grads: Optional[torch.Tensor] = None
        self._handles: list = []

    def __enter__(self) -> "GradCAM":
        self._handles.append(
            self.target_layer.register_forward_hook(self._fwd_hook)
        )
        self._handles.append(
            self.target_layer.register_full_backward_hook(self._bwd_hook)
        )
        return self

    def __exit__(self, *_) -> None:
        for h in self._handles:
            h.remove()
        self._handles.clear()

    def _fwd_hook(self, _m, _inp, out):
        self._acts = out[0] if isinstance(out, (tuple, list)) else out

    def _bwd_hook(self, _m, _gin, gout):
        self._grads = gout[0]

    def cam_from_score(self, score: torch.Tensor, retain_graph: bool = False) -> np.ndarray:
        self._grads = None
        self.model.zero_grad(set_to_none=True)
        score.backward(retain_graph=retain_graph)

        acts, grads = self._acts, self._grads
        if acts is None or grads is None:
            return np.zeros((7, 7), dtype=np.float32)

        weights = grads.mean(dim=(2, 3), keepdim=True)
        cam = torch.relu((weights * acts).sum(dim=1))[0]
        return cam.detach().cpu().numpy().astype(np.float32)


def compute_gradcam_for_layer(
    model: nn.Module,
    target_layer: nn.Module,
    face_tensor: torch.Tensor,
    device: torch.device,
    topk_bins: int = 8,
    mode: str = "both",
) -> Tuple[np.ndarray, np.ndarray]:
    """
    Reformed Grad-CAM: one forward pass, at most two backward passes.

    Instead of looping over N_bins individual backward passes, each head's
    contribution is collapsed to a single probability-weighted score before
    backprop, making it O(1) passes per head regardless of topk_bins.

    Returns (cam_yaw, cam_pitch), each float32 array normalized to [0, 1].
    Spatial size matches the target layer's feature map (e.g. 7x7 for layer4).
    """
    model.eval()
    x = face_tensor.detach().clone().to(device).float().requires_grad_(True)

    cam_yaw = np.zeros((7, 7), dtype=np.float32)
    cam_pitch = np.zeros((7, 7), dtype=np.float32)

    def _weighted_score(logits: torch.Tensor, probs: np.ndarray) -> torch.Tensor:
        K = min(topk_bins, probs.shape[0])
        top_idx = np.argsort(-probs)[:K]
        return sum(
            float(probs[i]) * logits[0, i]
            for i in top_idx
            if float(probs[i]) > 1e-8
        )

    with GradCAM(model, target_layer) as gcam:
        yaw_logits, pitch_logits = model(x)

        yaw_p = torch.softmax(yaw_logits.detach(), dim=1)[0].cpu().numpy()
        pitch_p = torch.softmax(pitch_logits.detach(), dim=1)[0].cpu().numpy()

        if mode == "sum":
            # Single backward through the combined yaw+pitch score.
            score = _weighted_score(yaw_logits, yaw_p) + _weighted_score(pitch_logits, pitch_p)
            cam = _norm01(gcam.cam_from_score(score, retain_graph=False))
            cam_yaw = cam_pitch = cam

        elif mode == "yaw":
            cam_yaw = _norm01(gcam.cam_from_score(
                _weighted_score(yaw_logits, yaw_p), retain_graph=False
            ))

        elif mode == "pitch":
            cam_pitch = _norm01(gcam.cam_from_score(
                _weighted_score(pitch_logits, pitch_p), retain_graph=False
            ))

        else:  # "both" — two separate backward passes
            cam_yaw = _norm01(gcam.cam_from_score(
                _weighted_score(yaw_logits, yaw_p), retain_graph=True
            ))
            cam_pitch = _norm01(gcam.cam_from_score(
                _weighted_score(pitch_logits, pitch_p), retain_graph=False
            ))

    return cam_yaw, cam_pitch


def save_cam_overlay(face_bgr_224: np.ndarray, cam01: np.ndarray, out_path: str, alpha: float = 0.6) -> None:
    H, W = face_bgr_224.shape[:2]
    cam_resized = cv2.resize(cam01, (W, H), interpolation=cv2.INTER_LINEAR)
    cam_u8 = (255.0 * cam_resized).clip(0, 255).astype(np.uint8)
    cam_color = cv2.applyColorMap(cam_u8, cv2.COLORMAP_MAGMA)
    out = (face_bgr_224.astype(np.float32) * (1 - alpha) + cam_color.astype(np.float32) * alpha)
    out = out.clip(0, 255).astype(np.uint8)
    cv2.imwrite(out_path, out)


# =============================================================================
# SINGLE OBJECT EXTRACTION
# =============================================================================
def extract_single_object_heatmap(
    H: np.ndarray,
    rel_thresh: float = 0.5,
    smooth_sigma: float = 1.2,
    min_area: int = 30,
    use_morph_close: bool = True,
) -> Tuple[np.ndarray, Tuple[int, int], np.ndarray]:
    H = H.astype(np.float32)

    if smooth_sigma and smooth_sigma > 0:
        Hs = cv2.GaussianBlur(H, ksize=(0, 0), sigmaX=float(smooth_sigma), sigmaY=float(smooth_sigma))
    else:
        Hs = H

    peak_idx = int(np.argmax(Hs))
    py, px = np.unravel_index(peak_idx, Hs.shape)
    peak_val = float(Hs[py, px])

    if not np.isfinite(peak_val) or peak_val <= 0:
        H_obj = np.full_like(H, 1.0 / H.size, dtype=np.float32)
        region_mask = np.ones_like(H, dtype=bool)
        return H_obj, (int(py), int(px)), region_mask

    thr = float(rel_thresh) * peak_val
    mask = (Hs >= thr).astype(np.uint8)

    if use_morph_close:
        k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (5, 5))
        mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, k, iterations=1)

    num_labels, labels, _stats, _ = cv2.connectedComponentsWithStats(mask, connectivity=8)

    if num_labels <= 1:
        region = mask.astype(bool)
    else:
        peak_label = int(labels[py, px])
        region = labels == peak_label

    if int(region.sum()) < int(min_area):
        thr2 = 0.25 * peak_val
        mask2 = (Hs >= thr2).astype(np.uint8)
        if use_morph_close:
            k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (5, 5))
            mask2 = cv2.morphologyEx(mask2, cv2.MORPH_CLOSE, k, iterations=1)
        num_labels2, labels2, _stats2, _2 = cv2.connectedComponentsWithStats(mask2, connectivity=8)
        if num_labels2 > 1:
            peak_label2 = int(labels2[py, px])
            region2 = labels2 == peak_label2
            if int(region2.sum()) >= int(min_area):
                region = region2

    H_obj = H * region.astype(np.float32)
    s = float(H_obj.sum())
    if np.isfinite(s) and s > 0:
        H_obj /= s
    else:
        H_obj[:] = 0.0
        H_obj[py, px] = 1.0

    return H_obj.astype(np.float32), (int(py), int(px)), region.astype(bool)


# =============================================================================
# METRICS (weighted AUC/AP with soft GT)
# =============================================================================
def _soft_gt_sigma_px(Hd: int, Wd: int) -> float:
    if GT_SOFT_SIGMA_MODE == "pixels":
        sigma = float(GT_SOFT_SIGMA_VALUE)
    else:
        sigma = float(GT_SOFT_SIGMA_VALUE) * float(min(Hd, Wd))
    return float(np.clip(sigma, GT_SOFT_SIGMA_MIN_PX, GT_SOFT_SIGMA_MAX_PX))


def _make_soft_gt(gt_mask_u8: np.ndarray, Hd: int, Wd: int) -> np.ndarray:
    gt = (gt_mask_u8 > 0).astype(np.float32)
    if gt.sum() <= 0:
        return gt

    sigma = _soft_gt_sigma_px(Hd, Wd)
    if sigma > 0:
        gt = cv2.GaussianBlur(gt, ksize=(0, 0), sigmaX=sigma, sigmaY=sigma)

    mx = float(gt.max())
    if mx > 0:
        gt = gt / mx
    return gt.astype(np.float32)


def _weighted_auc(scores: np.ndarray, pos_w: np.ndarray, eps: float = 1e-12) -> Optional[float]:
    s = scores.astype(np.float64).ravel()
    w_pos = pos_w.astype(np.float64).ravel()
    w_neg = 1.0 - w_pos

    Wp = float(w_pos.sum())
    Wn = float(w_neg.sum())
    if Wp <= eps or Wn <= eps:
        return None

    order = np.argsort(s)
    s_sorted = s[order]
    w_pos_sorted = w_pos[order]
    w_neg_sorted = w_neg[order]

    cneg = np.cumsum(w_neg_sorted)

    auc_num = 0.0
    i = 0
    N = s_sorted.size
    while i < N:
        j = i
        while j + 1 < N and s_sorted[j + 1] == s_sorted[i]:
            j += 1

        cneg_before = cneg[i - 1] if i > 0 else 0.0
        cneg_at_j = cneg[j]
        cneg_mid = 0.5 * (cneg_before + cneg_at_j)

        wpos_block = float(w_pos_sorted[i : j + 1].sum())
        auc_num += wpos_block * cneg_mid

        i = j + 1

    return float(auc_num / (Wp * Wn))


def _weighted_average_precision(scores: np.ndarray, pos_w: np.ndarray, eps: float = 1e-12) -> Optional[float]:
    s = scores.astype(np.float64).ravel()
    w_pos = pos_w.astype(np.float64).ravel()
    w_neg = 1.0 - w_pos

    P = float(w_pos.sum())
    if P <= eps:
        return None

    order = np.argsort(-s)
    w_pos = w_pos[order]
    w_neg = w_neg[order]

    tp = np.cumsum(w_pos)
    fp = np.cumsum(w_neg)

    precision = tp / np.maximum(tp + fp, eps)
    recall = tp / P

    ap = 0.0
    prev_r = 0.0
    for p, r in zip(precision, recall):
        ap += float(p) * float(r - prev_r)
        prev_r = float(r)
    return float(ap)


def _mask_centroid(mask_u8: np.ndarray) -> Optional[Tuple[float, float]]:
    ys, xs = np.where(mask_u8 > 0)
    if xs.size == 0:
        return None
    return (float(xs.mean()), float(ys.mean()))


def _angle_deg_between(a: np.ndarray, b: np.ndarray) -> float:
    a = a.astype(np.float64)
    b = b.astype(np.float64)
    a = a / (np.linalg.norm(a) + 1e-12)
    b = b / (np.linalg.norm(b) + 1e-12)
    c = float(np.clip(np.dot(a, b), -1.0, 1.0))
    return float(np.rad2deg(np.arccos(c)))


# =============================================================================
# GROUND TRUTH ANNOTATION
# =============================================================================
def _poly_to_mask(full_h: int, full_w: int, poly_xy: List[List[int]]) -> np.ndarray:
    mask = np.zeros((full_h, full_w), dtype=np.uint8)
    if poly_xy is None or len(poly_xy) < 3:
        return mask
    pts = np.array(poly_xy, dtype=np.int32).reshape(-1, 1, 2)
    cv2.fillPoly(mask, [pts], 255)
    return mask


def _annotate_polygon_on_image(ref_bgr: np.ndarray, win_name: str) -> Optional[List[List[int]]]:
    points: List[Tuple[int, int]] = []
    saved: Optional[List[List[int]]] = None
    disp = ref_bgr.copy()

    def draw_overlay(img: np.ndarray) -> np.ndarray:
        out = img.copy()
        panel_h = 150
        panel_w = 620
        cv2.rectangle(out, (5, 5), (5 + panel_w, 5 + panel_h), (0, 0, 0), -1)
        cv2.rectangle(out, (5, 5), (5 + panel_w, 5 + panel_h), (255, 255, 255), 2)
        lines = [
            "ANNOTATION (polygon on this reference image)",
            "Left click : add point",
            "Right click: finish + SAVE (>=3 points)",
            "u: undo  r: reset  s: save  q/ESC: quit",
            f"Points: {len(points)}",
        ]
        y = 35
        for i, t in enumerate(lines):
            fs = 0.75 if i == 0 else 0.65
            th = 2
            cv2.putText(out, t, (15, y), cv2.FONT_HERSHEY_SIMPLEX, fs, (255, 255, 255), th, cv2.LINE_AA)
            y += 26

        for i, (x, y_) in enumerate(points):
            cv2.circle(out, (x, y_), 5, (0, 0, 255), -1)
            cv2.putText(out, str(i + 1), (x + 6, y_ - 6), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 255), 2, cv2.LINE_AA)
            if i > 0:
                cv2.line(out, points[i - 1], points[i], (0, 255, 0), 2)
        if len(points) >= 3:
            cv2.line(out, points[-1], points[0], (0, 255, 0), 1)
        return out

    def redraw():
        nonlocal disp
        disp = draw_overlay(ref_bgr)

    def on_mouse(event, x, y, flags, param):
        nonlocal points, saved
        if event == cv2.EVENT_LBUTTONDOWN:
            points.append((int(x), int(y)))
            redraw()
        elif event == cv2.EVENT_RBUTTONDOWN:
            if len(points) >= 3:
                saved = [[int(px), int(py)] for (px, py) in points]

    cv2.namedWindow(win_name, cv2.WINDOW_NORMAL)
    cv2.setMouseCallback(win_name, on_mouse)
    redraw()

    while True:
        cv2.imshow(win_name, disp)
        k = cv2.waitKey(20) & 0xFF

        if saved is not None:
            break

        if k == ord("u") and len(points) > 0:
            points.pop()
            redraw()
        elif k == ord("r"):
            points = []
            redraw()
        elif k == ord("s"):
            if len(points) >= 3:
                saved = [[int(px), int(py)] for (px, py) in points]
                break
        elif k == ord("q") or k == 27:
            saved = None
            break

    cv2.destroyWindow(win_name)
    return saved


def _make_contact_sheet(primary_paths: List[str], max_imgs: int = 6, thumb_w: int = 480) -> np.ndarray:
    sel = primary_paths[:max_imgs]
    thumbs = []
    for p in sel:
        bgr = np.array(Image.open(p).convert("RGB"))[:, :, ::-1]
        h, w = bgr.shape[:2]
        scale = thumb_w / float(w)
        th = int(round(h * scale))
        t = cv2.resize(bgr, (thumb_w, th), interpolation=cv2.INTER_AREA)

        bar_h = 34
        bar = np.zeros((bar_h, thumb_w, 3), dtype=np.uint8)
        name = os.path.basename(p)
        cv2.putText(bar, name, (10, 24), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2, cv2.LINE_AA)
        t = np.vstack([bar, t])
        thumbs.append(t)

    if len(thumbs) == 0:
        return np.zeros((200, 600, 3), dtype=np.uint8)

    def pad_to_h(img, H):
        if img.shape[0] >= H:
            return img
        pad = np.zeros((H - img.shape[0], img.shape[1], 3), dtype=np.uint8)
        return np.vstack([img, pad])

    cols = 3
    rows = int(np.ceil(len(thumbs) / cols))

    row_imgs = []
    for r in range(rows):
        chunk = thumbs[r * cols : (r + 1) * cols]
        Hmax = max(im.shape[0] for im in chunk)
        chunk = [pad_to_h(im, Hmax) for im in chunk]
        while len(chunk) < cols:
            chunk.append(np.zeros((Hmax, thumb_w, 3), dtype=np.uint8))
        row_imgs.append(np.hstack(chunk))

    sheet = np.vstack(row_imgs)
    return sheet


def ensure_ground_truth_for_scenario(
    scenario_name: str,
    scenario_dir: str,
    scenario_out_dir: str,
    ref_paths: List[str],
    deepgaze_view_idxs: List[int],
    primary_paths: List[str],
    gt_root: str,
) -> Dict[str, Dict[str, List[List[int]]]]:
    gt_old_path = os.path.join(gt_root, scenario_name, GT_FILENAME)
    if os.path.isfile(gt_old_path):
        with open(gt_old_path, "r") as f:
            return json.load(f)

    gt_new_path = os.path.join(scenario_out_dir, GT_FILENAME)
    if os.path.isfile(gt_new_path):
        with open(gt_new_path, "r") as f:
            return json.load(f)

    if not ANNOTATE_IF_MISSING:
        return {}

    os.makedirs(scenario_out_dir, exist_ok=True)

    sheet = _make_contact_sheet(primary_paths, max_imgs=6, thumb_w=480)
    header_h = 80
    header = np.zeros((header_h, sheet.shape[1], 3), dtype=np.uint8)
    cv2.putText(header, f"SCENARIO: {scenario_name}", (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (255, 255, 255), 2, cv2.LINE_AA)
    cv2.putText(header, f"FOLDER: {scenario_dir}", (10, 65), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2, cv2.LINE_AA)
    context = np.vstack([header, sheet])

    cv2.namedWindow("SCENARIO CONTEXT (primaries)", cv2.WINDOW_NORMAL)
    cv2.imshow("SCENARIO CONTEXT (primaries)", context)
    cv2.waitKey(50)

    gt: Dict[str, Dict[str, List[List[int]]]] = {}

    for v in deepgaze_view_idxs:
        ref_index = v - 1
        if ref_index < 0 or ref_index >= len(ref_paths):
            continue

        ref_bgr = np.array(Image.open(ref_paths[ref_index]).convert("RGB"))[:, :, ::-1].copy()

        overlay = ref_bgr.copy()
        cv2.putText(overlay, f"{scenario_name}", (10, 30), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (255, 255, 255), 2, cv2.LINE_AA)
        cv2.putText(
            overlay,
            f"{os.path.basename(ref_paths[ref_index])} (view idx {v})",
            (10, 60),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.8,
            (255, 255, 255),
            2,
            cv2.LINE_AA,
        )

        win = f"Annotate GT polygon: {scenario_name} (view idx {v})"
        poly = _annotate_polygon_on_image(overlay, win_name=win)

        if poly is not None:
            gt[f"view_{v}"] = {
                "polygon_fullres": poly,
                "ref_path": ref_paths[ref_index],
            }

    cv2.destroyWindow("SCENARIO CONTEXT (primaries)")

    with open(gt_new_path, "w") as f:
        json.dump(gt, f, indent=2)

    try:
        os.makedirs(os.path.dirname(gt_old_path), exist_ok=True)
        with open(gt_old_path, "w") as f:
            json.dump(gt, f, indent=2)
    except Exception:
        pass

    return gt


# =============================================================================
# INPUT DISCOVERY + RESUME HELPERS
# =============================================================================
def list_images_in_dir(d: str) -> List[str]:
    exts = ["*.png", "*.jpg", "*.jpeg", "*.PNG", "*.JPG", "*.JPEG"]
    paths: List[str] = []
    for e in exts:
        paths += glob(os.path.join(d, e))
    return sorted(paths)


def discover_scenario_roots(base: str) -> List[str]:
    """
    Walk base recursively and return every folder that contains all REF_IMAGE_NAMES.
    That folder is a "data root" (has view1.jpg + view2.jpg) and its sub-folders
    are the individual scenarios to process.
    Once a data root is found, its subtree is not searched further.
    """
    ref_names_lower = {n.lower() for n in REF_IMAGE_NAMES}
    roots: List[str] = []
    for dirpath, dirnames, filenames in os.walk(base):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith("."))
        if ref_names_lower.issubset({f.lower() for f in filenames}):
            roots.append(dirpath)
            dirnames.clear()  # don't descend into a data root
    return sorted(roots)


def list_scenario_dirs(input_root: str, output_root: str) -> List[str]:
    """Return sub-scenario folders inside input_root, excluding output_root."""
    dirs = []
    for p in sorted(glob(os.path.join(input_root, "*"))):
        if not os.path.isdir(p):
            continue
        if os.path.basename(p).startswith("."):
            continue
        if os.path.abspath(p) == os.path.abspath(output_root):
            continue
        # only include folders that contain images
        if len(list_images_in_dir(p)) > 0:
            dirs.append(p)
    return dirs


def _marker_path(out_dir: str, which: str) -> str:
    return os.path.join(out_dir, which)


def _already_done(out_dir: str) -> bool:
    return os.path.isfile(_marker_path(out_dir, DONE_MARKER))


def _already_skipped(out_dir: str) -> bool:
    return os.path.isfile(_marker_path(out_dir, SKIP_MARKER))


def _write_marker(out_dir: str, which: str, payload: Dict[str, Any]) -> None:
    try:
        with open(_marker_path(out_dir, which), "w") as f:
            json.dump(payload, f, indent=2)
    except Exception:
        pass


# =============================================================================
# CORE RUN — split into three parts for ablation efficiency
# =============================================================================

def _run_shared_features(
    primary_path: str,
    ref_paths: List[str],
    out_dir_base: str,
    vggt_model: VGGT,
    l2cs_pipeline: Pipeline,
) -> Optional[Dict]:
    """
    Run VGGT + L2CS + Grad-CAM (shared across all cone conditions).
    Returns a dict of shared results, or None on fatal error.
    Face crops and Grad-CAM images are saved directly into out_dir_base.
    """
    image_paths = [primary_path] + ref_paths

    # VGGT
    try:
        orig_images: List[np.ndarray] = []
        orig_sizes: List[Tuple[int, int]] = []
        for p in image_paths:
            arr = np.array(Image.open(p).convert("RGB"))
            orig_images.append(arr)
            orig_sizes.append((arr.shape[0], arr.shape[1]))

        images = load_and_preprocess_images(image_paths).to(device_str)
        _S, _C, H_net, W_net = images.shape

        with torch.inference_mode():
            with torch.cuda.amp.autocast(enabled=(device_str == "cuda"), dtype=dtype):
                predictions = vggt_model(images)
            extrinsic, intrinsic = pose_encoding_to_extri_intri(predictions["pose_enc"], images.shape[-2:])

        world_points = predictions["world_points"].detach().cpu().numpy()[0]
        extrinsic_np = extrinsic.detach().cpu().numpy()[0]
        intrinsic_np = intrinsic.detach().cpu().numpy()[0]

        del images, predictions, extrinsic, intrinsic
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

    except Exception as e:
        _write_marker(out_dir_base, SKIP_MARKER, {"primary": primary_path, "status": "vggt_failed", "error": str(e)})
        return None

    # L2CS detection
    gaze_img = orig_images[GAZE_VIEW_IDX][:, :, ::-1].copy()
    try:
        res = l2cs_pipeline.step(gaze_img)
    except Exception as e:
        _write_marker(out_dir_base, SKIP_MARKER, {"primary": primary_path, "status": "l2cs_step_failed", "error": str(e)})
        return None

    if res.bboxes is None or len(res.bboxes) == 0:
        _write_marker(out_dir_base, SKIP_MARKER, {"primary": primary_path, "status": "no_face"})
        return None

    bbox = res.bboxes[0]

    # Face crops (saved once, shared by all cone conditions)
    h0, w0 = gaze_img.shape[:2]
    x1, y1, x2, y2 = _clip_bbox_xyxy(np.array(bbox, dtype=np.float32), h0, w0)
    cv2.imwrite(os.path.join(out_dir_base, "face_crop_raw.png"), gaze_img[y1:y2, x1:x2].copy())

    face_t, face_disp_bgr_224 = preprocess_face_for_l2cs_with_vis(
        gaze_img, np.array(bbox, dtype=np.float32), torch_dev, out_size=224,
    )
    cv2.imwrite(os.path.join(out_dir_base, "face_crop_224.png"), face_disp_bgr_224)

    # Eye origin
    origin3d, _origin_px = estimate_eye_origin_world(
        bbox=bbox,
        world_points_view=world_points[GAZE_VIEW_IDX],
        orig_size=orig_sizes[GAZE_VIEW_IDX],
        H_net=H_net, W_net=W_net,
        intrinsic=intrinsic_np[GAZE_VIEW_IDX],
        extrinsic=extrinsic_np[GAZE_VIEW_IDX],
    )

    # L2CS forward
    try:
        with torch.inference_mode():
            yaw_logits, pitch_logits = l2cs_pipeline.model(face_t)
    except Exception as e:
        _write_marker(out_dir_base, SKIP_MARKER, {"primary": primary_path, "status": "l2cs_forward_failed", "error": str(e)})
        return None

    yaw_probs = torch.softmax(yaw_logits, dim=1).detach().cpu().numpy()[0]
    pitch_probs = torch.softmax(pitch_logits, dim=1).detach().cpu().numpy()[0]

    bin_angles_deg = np.arange(NUM_BINS, dtype=np.float32) * BIN_WIDTH_DEG + BIN_START_DEG
    yaw_deg = expected_angle_deg(yaw_probs, bin_angles_deg)
    pitch_deg = expected_angle_deg(pitch_probs, bin_angles_deg)

    # Grad-CAM (saved once into out_dir_base, same for all cone conditions)
    if USE_GRADCAM_L2CS:
        cudnn_prev = torch.backends.cudnn.enabled
        torch.backends.cudnn.enabled = False
        try:
            cam_device = torch.device("cpu") if GRADCAM_FORCE_CPU else torch_dev
            l2cs_pipeline.model = l2cs_pipeline.model.to(cam_device).float()
            l2cs_pipeline.model.eval()
            layers_dict = find_resnet_layers(l2cs_pipeline.model)
            target_layers = (
                {GRADCAM_SINGLE_LAYER_NAME: layers_dict.get(GRADCAM_SINGLE_LAYER_NAME, layers_dict["lastconv"])}
                if GRADCAM_LAYER_MODE == "single" else layers_dict
            )
            with open(os.path.join(out_dir_base, "gradcam_layer_info.txt"), "w") as f:
                f.write(f"gradcam_device={cam_device}\nmethod=weighted_score_single_backward_per_head\n"
                        f"topk_bins={GRADCAM_TOPK_BINS}\nlayer_mode={GRADCAM_LAYER_MODE}\n"
                        f"layers_used={list(target_layers.keys())}\ncompute_mode={GRADCAM_COMPUTE}\n"
                        "torch.backends.cudnn.enabled=False (Grad-CAM only)\n")
            face_t_cam = face_t.detach().clone().to(cam_device).float()
            for layer_name, layer_mod in target_layers.items():
                cam_yaw, cam_pitch = compute_gradcam_for_layer(
                    l2cs_pipeline.model, layer_mod, face_t_cam, cam_device,
                    topk_bins=GRADCAM_TOPK_BINS, mode=GRADCAM_COMPUTE,
                )
                if GRADCAM_COMPUTE in ("yaw", "both", "sum"):
                    save_cam_overlay(face_disp_bgr_224, cam_yaw,
                                     os.path.join(out_dir_base, f"gradcam_yaw_{layer_name}.png"), alpha=GRADCAM_ALPHA)
                if GRADCAM_COMPUTE in ("pitch", "both", "sum"):
                    save_cam_overlay(face_disp_bgr_224, cam_pitch,
                                     os.path.join(out_dir_base, f"gradcam_pitch_{layer_name}.png"), alpha=GRADCAM_ALPHA)
                if SAVE_GRADCAM_NPY:
                    np.save(os.path.join(out_dir_base, f"gradcam_yaw_{layer_name}.npy"), cam_yaw)
                    np.save(os.path.join(out_dir_base, f"gradcam_pitch_{layer_name}.npy"), cam_pitch)
                # Free computation graph memory after each layer — 5 layers × 2 backward
                # passes accumulate significant RAM on CPU without this.
                del cam_yaw, cam_pitch
                l2cs_pipeline.model.zero_grad(set_to_none=True)
                gc.collect()
            del face_t_cam
        except Exception as e:
            with open(os.path.join(out_dir_base, "gradcam_error.txt"), "w") as f:
                f.write(str(e))
        finally:
            torch.backends.cudnn.enabled = cudnn_prev
            l2cs_pipeline.model = l2cs_pipeline.model.to(torch_dev)
            l2cs_pipeline.model.zero_grad(set_to_none=True)
            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()

    # Gaze direction in world (shared)
    dir_cam_gaze = gaze_vector_from_angles(float(np.deg2rad(yaw_deg)), float(np.deg2rad(pitch_deg)))
    dir_cam_gaze /= (np.linalg.norm(dir_cam_gaze) + 1e-9)
    dir_world = (extrinsic_np[GAZE_VIEW_IDX][:3, :3].T @ dir_cam_gaze).astype(np.float32)
    dir_world /= (np.linalg.norm(dir_world) + 1e-9)

    return {
        "orig_images": orig_images,
        "orig_sizes": orig_sizes,
        "world_points": world_points,
        "extrinsic_np": extrinsic_np,
        "intrinsic_np": intrinsic_np,
        "H_net": H_net,
        "W_net": W_net,
        "origin3d": origin3d,
        "yaw_probs": yaw_probs,
        "pitch_probs": pitch_probs,
        "bin_angles_deg": bin_angles_deg,
        "yaw_deg": yaw_deg,
        "pitch_deg": pitch_deg,
        "dir_world": dir_world,
    }


def _run_cone_condition(
    shared: Dict,
    cone_cfg: Dict,
    primary_path: str,
    out_dir: str,
    dg_model,
    centerbias_template: np.ndarray,
    gt_for_scenario: Dict[str, Dict[str, List[List[int]]]],
    metrics_rows: List[Dict],
) -> None:
    """Run DeepGaze + fusion + eval for one cone configuration."""
    os.makedirs(out_dir, exist_ok=True)

    yaw_probs    = shared["yaw_probs"]
    pitch_probs  = shared["pitch_probs"]
    bin_angles_deg = shared["bin_angles_deg"]
    yaw_deg      = shared["yaw_deg"]
    pitch_deg    = shared["pitch_deg"]
    dir_world    = shared["dir_world"]
    origin3d     = shared["origin3d"]
    world_points = shared["world_points"]
    orig_images  = shared["orig_images"]
    orig_sizes   = shared["orig_sizes"]
    extrinsic_np = shared["extrinsic_np"]
    intrinsic_np = shared["intrinsic_np"]
    H_net        = shared["H_net"]
    W_net        = shared["W_net"]

    # Cone angle for this condition
    if cone_cfg["dynamic"]:
        cone_angle_deg = cone_half_angle_from_probs(
            yaw_probs=yaw_probs, pitch_probs=pitch_probs,
            bin_angles_deg=bin_angles_deg,
            num_samples=CONE_SAMPLES, quantile=cone_cfg["quantile"],
            min_deg=CONE_MIN_DEG, max_deg=CONE_MAX_DEG,
        )
    else:
        cone_angle_deg = float(cone_cfg["fixed_deg"])

    cone_angle_rad = float(np.deg2rad(cone_angle_deg))
    sigma_theta = sigma_from_cone_quantile(cone_angle_rad, CONE_QUANTILE)

    pts_all = world_points.reshape(-1, 3)
    inside, _dist = compute_cone_mask_for_points(pts_all, origin3d, dir_world, cone_angle_deg)
    cone_pts = pts_all[inside] if inside.sum() > 0 else np.zeros((0, 3), dtype=np.float32)

    _projected_imgs, projected_masks = project_blob_to_images(
        blob_points=cone_pts,
        orig_images=orig_images, orig_sizes=orig_sizes,
        extrinsic_np=extrinsic_np, intrinsic_np=intrinsic_np,
        H_net=H_net, W_net=W_net, mask_method="kde",
    )

    with open(os.path.join(out_dir, "run_summary.txt"), "w") as f:
        f.write(f"primary={primary_path}\ncone_config={cone_cfg['name']}\n"
                f"dynamic={cone_cfg['dynamic']}\nyaw_deg={yaw_deg:.3f}\npitch_deg={pitch_deg:.3f}\n"
                f"cone_half_angle_deg={cone_angle_deg:.3f}\nsigma_theta_deg={np.rad2deg(sigma_theta):.3f}\n"
                f"use_gradcam={USE_GRADCAM_L2CS}\n")

    for v in DEEPGAZE_VIEW_IDXS:
        mask_u8 = _mask_to_u8(projected_masks[v])
        orig_bgr = orig_images[v][:, :, ::-1].copy()
        H0, W0 = orig_bgr.shape[:2]

        if mask_u8.shape[:2] != orig_bgr.shape[:2]:
            mask_u8 = cv2.resize(mask_u8, (W0, H0), interpolation=cv2.INTER_NEAREST)

        cv2.imwrite(os.path.join(out_dir, f"gaze_cone_mask_view_{v}.png"), mask_u8)

        cone_cond_bgr_full = blend_image_with_mask(orig_bgr, mask_u8, outside_dark=MASK_OUTSIDE_DARK)
        cv2.imwrite(os.path.join(out_dir, f"cone_conditioned_input_fullres_view_{v}.jpg"), cone_cond_bgr_full)

        image = cone_cond_bgr_full
        if max(H0, W0) > DEEPGAZE_MAX_DIM:
            scale = DEEPGAZE_MAX_DIM / float(max(H0, W0))
            image = cv2.resize(cone_cond_bgr_full, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)

        Hd, Wd = image.shape[:2]
        cv2.imwrite(os.path.join(out_dir, f"cone_conditioned_input_view_{v}.jpg"), image)

        centerbias = cv2.resize(centerbias_template, (Wd, Hd), interpolation=cv2.INTER_NEAREST).astype(np.float32)
        centerbias -= logsumexp(centerbias)

        image_tensor = torch.from_numpy(image.transpose(2, 0, 1)).unsqueeze(0).to(device_str)
        centerbias_tensor = torch.from_numpy(centerbias).unsqueeze(0).to(device_str)
        if device_str == "cuda":
            image_tensor = image_tensor.half()
            centerbias_tensor = centerbias_tensor.half()
        else:
            image_tensor = image_tensor.float()
            centerbias_tensor = centerbias_tensor.float()

        with torch.inference_mode():
            log_density_prediction = dg_model(image_tensor, centerbias_tensor)

        ld = log_density_prediction[0, 0].float().cpu().numpy().astype(np.float32)
        S_sal = np.exp(ld - float(ld.max())).astype(np.float32)
        S_sal = normalize_prob_map(S_sal)

        dir_cam_v = gaze_dir_cam_from_dir_world(dir_world, extrinsic_np[v])
        theta_net = theta_map_from_eye_and_worldpoints(
            world_points_view=world_points[v], extrinsic_w2c=extrinsic_np[v],
            eye_world=origin3d, gaze_dir_cam=dir_cam_v,
        )
        theta = cv2.resize(theta_net, (Wd, Hd), interpolation=cv2.INTER_LINEAR).astype(np.float32)

        alpha_val = float(ALPHA) if ALPHA is not None else float(np.clip(0.05 / (sigma_theta ** 2 + 1e-12), 0.5, 6.0))
        G = normalize_prob_map(geometry_likelihood(theta, sigma_theta))
        H_fused = fuse_poe(G, S_sal, alpha=alpha_val, beta=BETA, eps=EPS)

        np.save(os.path.join(out_dir, f"deepgaze_sal_view_{v}.npy"), S_sal)
        np.save(os.path.join(out_dir, f"theta_eyeorigin_rad_view_{v}.npy"), theta)
        np.save(os.path.join(out_dir, f"geom_G_view_{v}.npy"), G)
        np.save(os.path.join(out_dir, f"fused_H_view_{v}.npy"), H_fused)

        overlay_heatmap_bgr(image, S_sal, os.path.join(out_dir, f"_deepgaze_overlay_view_{v}.png"), alpha=0.6)
        overlay_heatmap_bgr(image, G,     os.path.join(out_dir, f"_geometry_overlay_view_{v}.png"), alpha=0.6)
        overlay_heatmap_bgr(image, H_fused, os.path.join(out_dir, f"_fused_overlay_view_{v}.png"), alpha=0.6)

        H_used = H_fused
        peak_xy: Optional[Tuple[int, int]] = None
        if SINGLEOBJ_ENABLE:
            H_single, (py, px), _region = extract_single_object_heatmap(
                H_fused, rel_thresh=SINGLEOBJ_REL_THRESH,
                smooth_sigma=SINGLEOBJ_SMOOTH_SIGMA, min_area=SINGLEOBJ_MIN_AREA,
                use_morph_close=SINGLEOBJ_MORPH_CLOSE,
            )
            H_used = H_single
            peak_xy = (px, py)
            np.save(os.path.join(out_dir, f"fused_H_singleobj_view_{v}.npy"), H_single)
            overlay_heatmap_bgr(image, H_single, os.path.join(out_dir, f"_fused_singleobj_overlay_view_{v}.png"), alpha=0.6)

        # Evaluation
        if DO_EVAL and gt_for_scenario and (f"view_{v}" in gt_for_scenario):
            poly_full = gt_for_scenario[f"view_{v}"]["polygon_fullres"]
            gt_full_mask = _poly_to_mask(H0, W0, poly_full)
            gt_dg_mask = cv2.resize(gt_full_mask, (Wd, Hd), interpolation=cv2.INTER_NEAREST)
            scores = H_used.astype(np.float32)

            if USE_SOFT_GT_FOR_METRICS:
                gt_soft = _make_soft_gt(gt_dg_mask, Hd=Hd, Wd=Wd)
                auc = _weighted_auc(scores, gt_soft)
                ap = _weighted_average_precision(scores, gt_soft)
                gt_soft_u8 = (255.0 * gt_soft / (float(gt_soft.max()) + 1e-12)).astype(np.uint8)
                cv2.imwrite(os.path.join(out_dir, f"gt_soft_view_{v}.png"), gt_soft_u8)
            else:
                gt_labels = (gt_dg_mask > 0).astype(np.int32)
                auc = _weighted_auc(scores, gt_labels.astype(np.float32))
                ap = _weighted_average_precision(scores, gt_labels.astype(np.float32))

            if peak_xy is None:
                idx = int(np.argmax(scores))
                py, px = np.unravel_index(idx, scores.shape)
                peak_xy = (int(px), int(py))

            c = _mask_centroid(gt_dg_mask)
            dist_px = None if c is None else float(np.hypot(peak_xy[0] - c[0], peak_xy[1] - c[1]))

            sx = float(W_net) / float(orig_sizes[v][1])
            sy = float(H_net) / float(orig_sizes[v][0])
            poly_net = [[int(round(x * sx)), int(round(y * sy))] for x, y in poly_full]
            gt_net_mask = _poly_to_mask(H_net, W_net, poly_net)

            pts = world_points[v][gt_net_mask > 0].reshape(-1, 3)
            finite = np.isfinite(pts).all(axis=1)
            ang_err_l2cs = None
            ang_err_peak = None
            if finite.any():
                centroid3d = np.median(pts[finite], axis=0).astype(np.float32)
                dir_gt_world = (centroid3d - origin3d).astype(np.float32)
                if np.isfinite(dir_gt_world).all() and np.linalg.norm(dir_gt_world) > 1e-9:
                    # L2CS angular error: raw gaze direction vs GT
                    ang_err_l2cs = _angle_deg_between(dir_world, dir_gt_world)

                    # Method angular error: pipeline peak pixel → 3D world point → angle to GT
                    px_net = int(np.clip(round(peak_xy[0] * W_net / Wd), 0, W_net - 1))
                    py_net = int(np.clip(round(peak_xy[1] * H_net / Hd), 0, H_net - 1))
                    peak_world = world_points[v][py_net, px_net]
                    if np.isfinite(peak_world).all():
                        dir_peak_world = (peak_world - origin3d).astype(np.float32)
                        if np.linalg.norm(dir_peak_world) > 1e-9:
                            ang_err_peak = _angle_deg_between(dir_peak_world, dir_gt_world)

            # L2CS dist_px: where the L2CS gaze ray hits the image (minimum-theta pixel)
            dist_px_l2cs = None
            if c is not None:
                l2cs_py, l2cs_px = np.unravel_index(int(np.argmin(theta)), theta.shape)
                dist_px_l2cs = float(np.hypot(int(l2cs_px) - c[0], int(l2cs_py) - c[1]))

            metrics_rows.append({
                "primary_path": primary_path,
                "cone_config": cone_cfg["name"],
                "view_idx": v,
                "yaw_deg": float(yaw_deg),
                "pitch_deg": float(pitch_deg),
                "cone_half_angle_deg": float(cone_angle_deg),
                "sigma_theta_deg": float(np.rad2deg(sigma_theta)),
                "alpha": float(alpha_val),
                "beta": float(BETA),
                "auc": auc,
                "ap": ap,
                "peak_x": int(peak_xy[0]),
                "peak_y": int(peak_xy[1]),
                "dist_px_to_gt_centroid": dist_px,
                "dist_px_to_gt_centroid_l2cs": dist_px_l2cs,
                "angular_error_l2cs_deg": ang_err_l2cs,
                "angular_error_peak_deg": ang_err_peak,
                "gt_soft_sigma_px": float(_soft_gt_sigma_px(Hd, Wd)) if USE_SOFT_GT_FOR_METRICS else None,
                "use_gradcam": bool(USE_GRADCAM_L2CS),
            })

        del image_tensor, centerbias_tensor, log_density_prediction
        if torch.cuda.is_available():
            torch.cuda.empty_cache()

    _write_marker(out_dir, DONE_MARKER, {"primary": primary_path, "cone_config": cone_cfg["name"], "status": "done"})


def run_one_primary_ablation(
    primary_path: str,
    ref_paths: List[str],
    out_dir_base: str,
    vggt_model: VGGT,
    l2cs_pipeline: Pipeline,
    dg_model,
    centerbias_template: np.ndarray,
    gt_for_scenario: Dict[str, Dict[str, List[List[int]]]],
    metrics_rows: List[Dict],
) -> None:
    """
    Outer wrapper: run shared features once, then iterate over CONE_ABLATION.
    Each condition gets its own subdirectory: out_dir_base/<cone_cfg['name']>/
    """
    os.makedirs(out_dir_base, exist_ok=True)

    # Check if all conditions already done (full resume)
    if RESUME_AUTO and all(
        _already_done(os.path.join(out_dir_base, cfg["name"])) for cfg in CONE_ABLATION
    ):
        return

    shared = _run_shared_features(primary_path, ref_paths, out_dir_base, vggt_model, l2cs_pipeline)
    if shared is None:
        return  # fatal error already written to out_dir_base

    for cone_cfg in CONE_ABLATION:
        cone_out_dir = os.path.join(out_dir_base, cone_cfg["name"])
        if RESUME_AUTO and _already_done(cone_out_dir):
            print(f"      [cone-skip] {cone_cfg['name']} (DONE)")
            continue
        print(f"      [cone] {cone_cfg['name']}")
        try:
            _run_cone_condition(
                shared=shared,
                cone_cfg=cone_cfg,
                primary_path=primary_path,
                out_dir=cone_out_dir,
                dg_model=dg_model,
                centerbias_template=centerbias_template,
                gt_for_scenario=gt_for_scenario,
                metrics_rows=metrics_rows,
            )
        except Exception as e:
            _write_marker(cone_out_dir, SKIP_MARKER,
                          {"primary": primary_path, "cone_config": cone_cfg["name"],
                           "status": "exception", "error": str(e)})

    del shared
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()


# =============================================================================
# EXCEL SUMMARY (per-person per-distance, mirroring alpha=2 folder layout)
# =============================================================================

def write_person_distance_excel(
    person_label: str,
    distance: str,
    data_by_cone: Dict[str, Dict[str, List[Dict]]],
    out_dir: str,
) -> str:
    """
    Write one Excel file for a single person+distance combination.

    Layout:
      - One sheet per cone_config (ablation condition).
      - Within each sheet: sub-scenario blocks in the alpha=2 style —
          header row / data rows / bold average row / blank row / repeat.

    Parameters
    ----------
    data_by_cone : {cone_name: {subscenario_name: [row_dicts]}}
    """
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"{person_label} {distance}.xlsx")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    avg_fill    = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for cone_name, subscenario_map in sorted(data_by_cone.items()):
        ws = wb.create_sheet(title=cone_name[:31])  # Excel sheet name limit = 31 chars

        all_rows = [r for rows in subscenario_map.values() for r in rows]
        if not all_rows:
            continue
        fieldnames = list(all_rows[0].keys())

        excel_row = 1

        for subscenario_name, rows in sorted(subscenario_map.items()):
            if not rows:
                continue

            # -- Header row (repeated per block, like alpha=2) --
            for col_idx, fn in enumerate(fieldnames, 1):
                cell = ws.cell(row=excel_row, column=col_idx, value=fn)
                cell.font = Font(bold=True)
                cell.fill = header_fill
            excel_row += 1

            # -- Data rows --
            for row_dict in rows:
                for col_idx, fn in enumerate(fieldnames, 1):
                    raw = row_dict.get(fn, "")
                    try:
                        val: Any = float(raw) if raw not in ("", None, "None", "nan") else raw
                    except (ValueError, TypeError):
                        val = raw
                    ws.cell(row=excel_row, column=col_idx, value=val)
                excel_row += 1

            # -- Average summary row --
            avg_vals: Dict[str, Any] = {}
            for col in SUMMARY_AVG_COLS:
                if col not in fieldnames:
                    continue
                nums = []
                for r in rows:
                    raw = r.get(col, "")
                    try:
                        v = float(raw)
                        if np.isfinite(v):
                            nums.append(v)
                    except (ValueError, TypeError):
                        pass
                avg_vals[col] = float(np.mean(nums)) if nums else ""

            for col_idx, fn in enumerate(fieldnames, 1):
                val = avg_vals.get(fn, "")
                cell = ws.cell(row=excel_row, column=col_idx, value=val)
                cell.font = Font(bold=True)
                cell.fill = avg_fill
            excel_row += 1

            # -- Blank separator row --
            excel_row += 1

        # Auto-width columns (best-effort)
        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 40)

    wb.save(out_path)
    return out_path


# =============================================================================
# MAIN
# =============================================================================
def main() -> None:
    print(f"device: {device_str}, dtype: {dtype}")
    print(f"BASE_ROOT: {BASE_ROOT}")

    if not os.path.isfile(CENTERBIAS_NPY):
        raise RuntimeError(
            f"centerbias file not found: {CENTERBIAS_NPY}\n"
            f"Run from the directory that contains it, or set CENTERBIAS_NPY to an absolute path."
        )

    # ---- Discover all data roots (folders containing view1.jpg + view2.jpg) ----
    data_roots = discover_scenario_roots(BASE_ROOT)
    if not data_roots:
        raise RuntimeError(f"No folders with {REF_IMAGE_NAMES} found under: {BASE_ROOT}")

    if ONLY_SCENARIO_NAME is not None:
        data_roots = [r for r in data_roots if os.path.basename(r) == ONLY_SCENARIO_NAME]
        if not data_roots:
            raise RuntimeError(f"ONLY_SCENARIO_NAME='{ONLY_SCENARIO_NAME}' not found under {BASE_ROOT}")

    print(f"Found {len(data_roots)} scenario root(s):")
    for r in data_roots:
        print(f"  {r}")

    # ---- Load models once ----
    print("\n[LOAD] VGGT (once)")
    vggt_model = VGGT.from_pretrained("facebook/VGGT-1B").to(dtype).to(device_str)
    vggt_model.eval()

    print("[LOAD] L2CS (once)")
    l2cs_pipeline = Pipeline(weights=L2CS_WEIGHTS, arch="ResNet50", device=torch_dev)

    print("[LOAD] DeepGaze IIE (once)")
    dg_model = deepgaze_pytorch.DeepGazeIIE(pretrained=True).to(device_str)
    dg_model.eval()
    if device_str == "cuda":
        dg_model.half()

    print("[LOAD] centerbias (once)")
    centerbias_template = np.load(CENTERBIAS_NPY)

    # ---- Outer loop: one data root = one set of reference images ----
    for data_root in data_roots:
        output_root = os.path.join(data_root, "outputs_batch")
        gt_root = output_root
        os.makedirs(output_root, exist_ok=True)

        ref_paths = [os.path.join(data_root, n) for n in REF_IMAGE_NAMES]
        missing = [p for p in ref_paths if not os.path.isfile(p)]
        if missing:
            print(f"[SKIP data root] missing refs {missing} in {data_root}")
            continue

        scenario_dirs = list_scenario_dirs(data_root, output_root)
        if not scenario_dirs:
            print(f"[SKIP data root] no sub-scenario folders found in {data_root}")
            continue

        if ONLY_SUBSCENARIO_NAME is not None:
            scenario_dirs = [d for d in scenario_dirs if os.path.basename(d) == ONLY_SUBSCENARIO_NAME]
            if not scenario_dirs:
                print(f"[SKIP data root] ONLY_SUBSCENARIO_NAME='{ONLY_SUBSCENARIO_NAME}' not found in {data_root}")
                continue

        print(f"\n{'='*80}")
        print(f"[DATA ROOT] {data_root}")
        print(f"  refs       : {ref_paths}")
        print(f"  output_root: {output_root}")
        print(f"  sub-scenarios ({len(scenario_dirs)}): {[os.path.basename(d) for d in scenario_dirs]}")
        print(f"{'='*80}")

        # ---- Inner loop: one sub-scenario = one set of primary images ----
        for scen_dir in scenario_dirs:
            scenario_name = os.path.basename(scen_dir)
            scenario_out_dir = os.path.join(output_root, scenario_name)
            os.makedirs(scenario_out_dir, exist_ok=True)

            primary_list = list_images_in_dir(scen_dir)
            if not primary_list:
                print(f"[SKIP] '{scenario_name}' has no images: {scen_dir}")
                continue

            print(f"\n[SCENARIO] {scenario_name}  ({len(primary_list)} primaries)")

            gt_for_scenario: Dict[str, Dict[str, List[List[int]]]] = {}
            if DO_EVAL:
                gt_for_scenario = ensure_ground_truth_for_scenario(
                    scenario_name=scenario_name,
                    scenario_dir=scen_dir,
                    scenario_out_dir=scenario_out_dir,
                    ref_paths=ref_paths,
                    deepgaze_view_idxs=DEEPGAZE_VIEW_IDXS,
                    primary_paths=primary_list,
                    gt_root=gt_root,
                )

            metrics_rows: List[Dict] = []

            for i, primary_path in enumerate(primary_list):
                base = os.path.splitext(os.path.basename(primary_path))[0]
                out_dir = os.path.join(scenario_out_dir, base)
                os.makedirs(out_dir, exist_ok=True)

                # Per-condition done checks happen inside run_one_primary_ablation.
                # Only skip here if the primary was globally skipped (VGGT/L2CS failure).
                if RESUME_AUTO and RESUME_SKIP_SKIPPED and _already_skipped(out_dir):
                    print(f"  [SKIP] {i+1}/{len(primary_list)} {os.path.basename(primary_path)} (SKIPPED)")
                    continue

                print(f"  [BATCH] {i+1}/{len(primary_list)} {os.path.basename(primary_path)}")
                run_one_primary_ablation(
                    primary_path=primary_path,
                    ref_paths=ref_paths,
                    out_dir_base=out_dir,
                    vggt_model=vggt_model,
                    l2cs_pipeline=l2cs_pipeline,
                    dg_model=dg_model,
                    centerbias_template=centerbias_template,
                    gt_for_scenario=gt_for_scenario,
                    metrics_rows=metrics_rows,
                )

            if DO_EVAL and metrics_rows:
                csv_path = os.path.join(scenario_out_dir, "metrics.csv")
                keys = list(metrics_rows[0].keys())
                with open(csv_path, "w", newline="") as f:
                    w = csv.DictWriter(f, fieldnames=keys)
                    w.writeheader()
                    for r in metrics_rows:
                        w.writerow(r)
                print(f"  [SAVE] metrics: {csv_path}")

            gc.collect()
            if torch.cuda.is_available():
                torch.cuda.empty_cache()

    # ---- Per-person per-distance Excel summaries ----
    if DO_EVAL and SUMMARY_OUT_DIR:
        print("\n[SUMMARY] Building per-person per-distance Excel files...")

        # Gather all metrics.csv files: BASE_ROOT/person/scenario/outputs_batch/subscenario/metrics.csv
        metrics_files = sorted(glob(
            os.path.join(BASE_ROOT, "*", "*", "outputs_batch", "*", "metrics.csv")
        ))

        # Structure: {(person_label, distance): {cone_name: {subscenario: [rows]}}}
        summary_data: Dict[Tuple[str, str], Dict[str, Dict[str, List[Dict]]]] = {}

        for mf in metrics_files:
            parts = os.path.relpath(mf, BASE_ROOT).split(os.sep)
            # parts: [person_folder, scenario_folder, "outputs_batch", subscenario, "metrics.csv"]
            if len(parts) != 5:
                continue
            person_folder, scenario_folder, _, subscenario_name, _ = parts

            person_label = PERSON_LABEL_MAP.get(person_folder, person_folder.lower())
            m = re.search(r"_(\d+m)", scenario_folder)
            if not m:
                continue
            distance = m.group(1)

            try:
                with open(mf, newline="") as f:
                    rows = list(csv.DictReader(f))
            except Exception:
                continue
            if not rows:
                continue

            key = (person_label, distance)
            summary_data.setdefault(key, {})

            for row in rows:
                cone_name = row.get("cone_config", "default")
                summary_data[key].setdefault(cone_name, {})
                summary_data[key][cone_name].setdefault(subscenario_name, [])
                summary_data[key][cone_name][subscenario_name].append(row)

        os.makedirs(SUMMARY_OUT_DIR, exist_ok=True)
        for (person_label, distance), cone_data in sorted(summary_data.items()):
            out_path = write_person_distance_excel(
                person_label=person_label,
                distance=distance,
                data_by_cone=cone_data,
                out_dir=SUMMARY_OUT_DIR,
            )
            print(f"  [EXCEL] {out_path}")

    print("\nDONE")


if __name__ == "__main__":
    main()
